"""Convert this paper's IEEEtran LaTeX (limited subset) to an IEEE-styled two-column DOCX."""
import re, sys, os
from docx import Document
from docx.shared import Pt, Inches, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
FIGDIR = sys.argv[2]
OUT = sys.argv[3]
tex = open(SRC, encoding="utf-8").read()
body = tex.split("\\maketitle", 1)[1].split("\\end{document}")[0]

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
FONT = "Times New Roman"

# ------------------------------------------------------------------ label/cite numbering
cites = re.findall(r"\\bibitem\{([^}]+)\}", body)
cite_no = {k: i + 1 for i, k in enumerate(cites)}
labels = {}
sec_i, sub_i = 0, 0
fig_i, tab_i = 0, 0
for m in re.finditer(r"\\(section\*?|subsection)\{[^}]*\}|\\begin\{(figure\*?|table\*?)\}|\\label\{([^}]+)\}", body):
    if m.group(1) == "section":
        sec_i += 1; sub_i = 0; last = ROMAN[sec_i - 1]
    elif m.group(1) == "subsection":
        sub_i += 1; last = f"{ROMAN[sec_i - 1]}-{chr(64 + sub_i)}"
    elif m.group(1) == "section*":
        pass
    elif m.group(2):
        if m.group(2).startswith("figure"):
            fig_i += 1; last = str(fig_i)
        else:
            tab_i += 1; last = ROMAN[tab_i - 1]
    elif m.group(3):
        labels[m.group(3)] = last

# ------------------------------------------------------------------ inline conversion
MATH_CMDS = {r"\pm": "±", r"\times": "×", r"\sigma": "σ", r"\phi": "φ", r"\epsilon": "ε",
             r"\approx": "≈", r"\circ": "°", r"\geq": "≥", r"\leq": "≤", r"\,": " ", r"\min": "min",
             r"\max": "max", r"\rightarrow": "→", r"\cdot": "·"}


def clean_text(s):
    s = s.replace("~", "\u00a0").replace("---", "—").replace("--", "–")
    s = s.replace("``", "“").replace("''", "”").replace("`", "‘").replace("'", "’")
    s = s.replace(r"\%", "%").replace(r"\_", "_").replace(r"\&", "&").replace(r"\'{\i}", "í")
    s = s.replace(r'\"{u}', "ü").replace(r"\,", "\u2009").replace(r"\/", "")
    s = re.sub(r"\\cite\{([^}]+)\}", lambda m: "[" + ", ".join(str(cite_no[k.strip()]) for k in m.group(1).split(",")) + "]", s)
    s = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels.get(m.group(1), "?"), s)
    s = re.sub(r"\\label\{[^}]+\}", "", s)
    return s


def math_runs(m):
    """Return list of (text, italic, sub, sup) from a simple LaTeX math string."""
    m = m.replace(r"\boldsymbol", "").replace(r"\text", "")
    for k, v in MATH_CMDS.items():
        m = m.replace(k, v)
    out = []
    i = 0
    def grab(j):
        if j < len(m) and m[j] == "{":
            depth, k = 1, j + 1
            while depth:
                depth += {"{": 1, "}": -1}.get(m[k], 0); k += 1
            return m[j + 1:k - 1], k
        return m[j], j + 1
    while i < len(m):
        c = m[i]
        if c in "_^":
            t, i = grab(i + 1)
            t = t.replace("{", "").replace("}", "")
            out.append((t, True, c == "_", c == "^"))
        elif c in "{}":
            i += 1
        else:
            out.append((c, c.isalpha(), False, False)); i += 1
    return out


def add_runs(p, s, size=None, bold=False, italic=False, color=None):
    """Parse inline LaTeX (\\textit, \\emph, \\textbf, \\texttt, $math$) into runs."""
    s = clean_text(s)
    tokens = re.split(r"(\\textit\{|\\emph\{|\\textbf\{|\\texttt\{|\\url\{|\$[^$]*\$|\})", s)
    stack = []
    for tk in tokens:
        if not tk:
            continue
        if tk in ("\\textit{", "\\emph{", "\\textbf{", "\\texttt{", "\\url{"):
            stack.append(tk); continue
        if tk == "}":
            if stack: stack.pop()
            continue
        b = bold or "\\textbf{" in stack
        it = italic or "\\textit{" in stack or "\\emph{" in stack
        mono = "\\texttt{" in stack
        if tk.startswith("$") and tk.endswith("$") and len(tk) > 1:
            for t, mit, sub, sup in math_runs(tk[1:-1]):
                r = p.add_run(t)
                r.italic = mit or it; r.bold = b
                r.font.subscript = sub; r.font.superscript = sup
                if size: r.font.size = Pt(size)
            continue
        r = p.add_run(tk.replace("\\", ""))
        r.bold, r.italic = b, it
        if mono:
            r.font.name = "Courier New"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            r.font.size = Pt((size or 10) - 1)
        elif size:
            r.font.size = Pt(size)
        if color: r.font.color.rgb = color


# ------------------------------------------------------------------ document setup
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0

sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.625)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)
COLW = (8.5 - 1.25 - 0.25) / 2
TEXTW = 8.5 - 1.25


def set_cols(section, n):
    sp = section._sectPr
    for c in sp.findall(qn("w:cols")):
        sp.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n)); cols.set(qn("w:space"), str(int(0.25 * 1440)))
    sp.append(cols)


def new_section(n):
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    set_cols(s, n)
    return s


def para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.14, before=0, after=0, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Inches(first) if first else None
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.keep_with_next = keep
    return p


# title & author (one column)
title = re.search(r"\\title\{(.+?)\}\n", tex).group(1)
p = para(WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 10)
r = p.add_run(title); r.font.size = Pt(24)
p = para(WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 2); r = p.add_run("Jayprakash Bind"); r.font.size = Pt(11)
for line, it in [("Dept. of Electronics and Communication Engineering, Pranveer Singh Institute of Technology", True),
                 ("Kanpur, Uttar Pradesh, India", False),
                 ("JAVICS Automation and Decoration Pvt. Ltd.", True), ("Mirzapur, Uttar Pradesh, India", False),
                 ("jpbjayprakash@gmail.com", False)]:
    p = para(WD_ALIGN_PARAGRAPH.CENTER, 0)
    r = p.add_run(line); r.italic = it; r.font.size = Pt(10)
para(WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 10)
set_cols(sec, 1)
new_section(2)

# ------------------------------------------------------------------ body walk
abstract = re.search(r"\\begin\{abstract\}(.+?)\\end\{abstract\}", body, re.S).group(1).strip()
kw = re.search(r"\\begin\{IEEEkeywords\}(.+?)\\end\{IEEEkeywords\}", body, re.S).group(1).strip()
p = para(first=0.14, after=6)
r = p.add_run("Abstract"); r.bold = r.italic = True; r.font.size = Pt(9)
r = p.add_run("—"); r.bold = True; r.font.size = Pt(9)
add_runs(p, " ".join(abstract.split()), size=9, bold=True)
p = para(first=0.14, after=8)
r = p.add_run("Index Terms"); r.bold = r.italic = True; r.font.size = Pt(9)
r = p.add_run("—"); r.bold = True; r.font.size = Pt(9)
add_runs(p, " ".join(kw.split()), size=9, bold=True)

rest = body.split("\\end{IEEEkeywords}", 1)[1]
blocks = re.split(r"(\\begin\{(?:itemize|enumerate|table\*?|figure\*?|equation|thebibliography)\}.*?\\end\{(?:itemize|enumerate|table\*?|figure\*?|equation|thebibliography)\})", rest, flags=re.S)

sec_n, sub_n, fig_n, tab_n, eq_n = 0, 0, 0, 0, 0
first_para_after_intro = False


def heading1(t, numbered=True):
    global sec_n, sub_n
    if numbered:
        sec_n += 1; sub_n = 0
        t = f"{ROMAN[sec_n - 1]}. {t}"
    p = para(WD_ALIGN_PARAGRAPH.CENTER, 0, 8, 4, keep=True)
    r = p.add_run(clean_text(t)); r.font.small_caps = True


def heading2(t):
    global sub_n
    sub_n += 1
    p = para(WD_ALIGN_PARAGRAPH.LEFT, 0, 5, 3, keep=True)
    r = p.add_run(f"{chr(64 + sub_n)}. {clean_text(t)}"); r.italic = True


def text_block(chunk):
    global first_para_after_intro
    pos = 0
    for m in re.finditer(r"\\(section\*?|subsection)\{([^}]*)\}", chunk):
        emit_paras(chunk[pos:m.start()])
        if m.group(1) == "section":
            heading1(m.group(2))
        elif m.group(1) == "section*":
            heading1(m.group(2), numbered=False)
        else:
            heading2(m.group(2))
        pos = m.end()
    emit_paras(chunk[pos:])


def emit_paras(t):
    for para_txt in re.split(r"\n\s*\n", t):
        s = " ".join(para_txt.split())
        s = re.sub(r"\\label\{[^}]+\}", "", s).strip()
        if not s:
            continue
        m = re.match(r"\\IEEEPARstart\{(\w)\}\{(\w+)\}(.*)", s)
        p = para(first=0 if m else 0.14)
        if m:
            r = p.add_run(m.group(1)); r.font.size = Pt(20)
            r = p.add_run(m.group(2).upper()); r.font.size = Pt(10)
            add_runs(p, m.group(3))
        else:
            add_runs(p, s)


def list_block(b, numbered):
    items = re.split(r"\\item\s*", b.split("}", 1)[1].rsplit("\\end", 1)[0])[1:]
    for k, it in enumerate(items, 1):
        p = para(WD_ALIGN_PARAGRAPH.JUSTIFY, 0)
        p.paragraph_format.left_indent = Inches(0.25 if not numbered else 0.28)
        p.paragraph_format.first_line_indent = Inches(-0.15 if not numbered else -0.2)
        p.add_run(f"{k})\u00a0" if numbered else "•\u00a0")
        add_runs(p, " ".join(it.split()))


def shade_border(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side, val in (("top", top), ("bottom", bottom), ("left", None), ("right", None)):
        e = OxmlElement(f"w:{side}")
        if val:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(val)); e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)


def table_block(b, wide):
    global tab_n
    tab_n += 1
    cap = clean_text(re.search(r"\\caption\{(.+?)\}\n", b).group(1))
    if wide: new_section(1)
    p = para(WD_ALIGN_PARAGRAPH.CENTER, 0, 6, 0, keep=True); r = p.add_run(f"TABLE {ROMAN[tab_n - 1]}"); r.font.size = Pt(8)
    p = para(WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 3, keep=True)
    add_runs(p, cap, size=8)
    for rr in p.runs: rr.font.small_caps = True
    tab = re.search(r"\\toprule(.+?)\\bottomrule", b, re.S).group(1)
    tab = tab.replace("\\midrule", "")
    rows = [r.strip() for r in tab.split("\\\\") if r.strip()]
    rows = [[c.strip() for c in r.split("&")] for r in rows]
    # merge two-line header (second row starting with empty cells and bold)
    ncol = max(len(r) for r in rows)
    hdr_rows = 2 if rows[1][0] == "" and all(("textbf" in c or c == "") for c in rows[1]) else 1
    size = 8 if not wide else 8.5
    fr = [float(x) for x in re.findall(r"p\{([0-9.]+)\\columnwidth\}", b)]
    avail = TEXTW if wide else COLW
    if len(fr) == ncol:
        widths = [avail * f / sum(fr) * min(sum(fr), 0.98) for f in fr]
    elif wide:
        widths = [avail * 0.20] + [avail * 0.78 / (ncol - 1)] * (ncol - 1)
    else:
        widths = [avail * 0.30, avail * 0.17] + [avail * 0.50 / (ncol - 2)] * (ncol - 2)
    tot = sum(widths)
    widths = [max(w_, 0.42) for w_ in widths]
    widths = [w_ * tot / sum(widths) for w_ in widths]
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for ci, col in enumerate(t.columns):
        col.width = Inches(widths[ci])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            cells[ci].width = Inches(widths[ci])
            c = row[ci] if ci < len(row) else ""
            c = re.sub(r"\\multirow\{\d\}\{\*\}\{(.+)\}", r"\1", c)
            cp = cells[ci].paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or not wide and ncol <= 3 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(cp, c, size=size, bold=ri < hdr_rows)
            shade_border(cells[ci], top=8 if ri == 0 else None,
                         bottom=6 if ri == hdr_rows - 1 else (8 if ri == len(rows) - 1 else None))
    para(after=6)
    if wide: new_section(2)


def figure_block(b, wide):
    global fig_n
    fig_n += 1
    f = re.search(r"\\includegraphics\[[^]]*\]\{([^}]+)\}", b).group(1).replace(".pdf", ".png")
    cap = " ".join(re.search(r"\\caption\{(.+)\}\s*\n\\label", b, re.S).group(1).split())
    if wide: new_section(1)
    p = para(WD_ALIGN_PARAGRAPH.CENTER, 0, 6, 2, keep=True)
    w = TEXTW * (0.92 if "fig1" in f else 1.0) if wide else COLW * (0.86 if "fig2" in f else 1.0)
    p.add_run().add_picture(os.path.join(FIGDIR, f), width=Inches(w))
    p = para(WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 0, 8)
    r = p.add_run(f"Fig. {fig_n}.\u2002"); r.font.size = Pt(8)
    add_runs(p, cap, size=8)
    if wide: new_section(2)


def equation_block(b):
    global eq_n
    eq_n += 1
    m = b.split("}", 1)[1].rsplit("\\end", 1)[0].strip()
    p = para(WD_ALIGN_PARAGRAPH.LEFT, 0, 4, 4)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(COLW / 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(COLW), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    trail = ""
    if m.endswith(","):
        m, trail = m[:-1], ","
    add_runs(p, "$" + m + "$")
    p.add_run(trail + f"\t({eq_n})")


def bib_block(b):
    heading1("References", numbered=False)
    items = re.findall(r"\\bibitem\{([^}]+)\}(.+?)(?=\\bibitem|\\end\{thebibliography\})", b, re.S)
    for k, (key, txt) in enumerate(items, 1):
        p = para(WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 0, 1)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
        r = p.add_run(f"[{k}]\t"); r.font.size = Pt(8)
        add_runs(p, " ".join(txt.split()), size=8)


for b in blocks:
    if not b.strip():
        continue
    if b.startswith("\\begin{itemize}"):
        list_block(b, False)
    elif b.startswith("\\begin{enumerate}"):
        list_block(b, True)
    elif b.startswith("\\begin{table*}"):
        table_block(b, True)
    elif b.startswith("\\begin{table}"):
        table_block(b, False)
    elif b.startswith("\\begin{figure*}"):
        figure_block(b, True)
    elif b.startswith("\\begin{figure}"):
        figure_block(b, False)
    elif b.startswith("\\begin{equation}"):
        equation_block(b)
    elif b.startswith("\\begin{thebibliography}"):
        bib_block(b)
    else:
        text_block(b)

doc.core_properties.author = "Jayprakash Bind"
doc.core_properties.title = title
doc.save(OUT)
print("saved", OUT, "figs", fig_n, "tables", tab_n, "cites", len(cite_no))
